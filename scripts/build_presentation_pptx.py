#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Presentación Ejecutiva PowerPoint (.pptx) para MIGATO
Estilo Corporativo / Business / Gubernamental - 10 Láminas Estratégicas
Basado en INFORME_TESIS_PRESUPUESTO_MIGATO_V3
"""

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = Path(__file__).resolve().parent.parent
LOGO_PATH = PROJECT_ROOT / "assets" / "logo-migato.png"

# Paleta Corporativa / Business Ejecutiva
COLOR_BG = RGBColor(248, 250, 252)         # Slate 50 (#F8FAFC)
COLOR_WHITE = RGBColor(255, 255, 255)      # White (#FFFFFF)
COLOR_CARD_BG = RGBColor(255, 255, 255)    # White Card
COLOR_CARD_SUBTLE = RGBColor(241, 245, 249)# Slate 100 (#F1F5F9)
COLOR_BORDER = RGBColor(203, 213, 225)     # Slate 300 (#CBD5E1)
COLOR_BORDER_LIGHT = RGBColor(226, 232, 240)# Slate 200 (#E2E8F0)

COLOR_NAVY_DARK = RGBColor(15, 23, 42)     # Slate 900 (#0F172A)
COLOR_NAVY_CORP = RGBColor(30, 58, 138)    # Blue 900 / Corporate Navy (#1E3A8A)
COLOR_BLUE_MIGATO = RGBColor(2, 132, 199)  # Sky 600 (#0284C7)
COLOR_BLUE_LIGHT = RGBColor(224, 242, 254) # Sky 100 (#E0F2FE)

COLOR_TEXT_MAIN = RGBColor(15, 23, 42)     # Charcoal Slate (#0F172A)
COLOR_TEXT_BODY = RGBColor(51, 65, 85)     # Slate 700 (#334155)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500 (#64748B)

COLOR_SUCCESS_DARK = RGBColor(5, 150, 105) # Emerald 600 (#059669)
COLOR_SUCCESS_BG = RGBColor(236, 253, 245) # Emerald 50 (#ECFDF5)
COLOR_DANGER_DARK = RGBColor(220, 38, 38)  # Red 600 (#DC2626)
COLOR_DANGER_BG = RGBColor(254, 242, 242)  # Red 50 (#FEF2F2)

def create_deck(filename):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="MIGATO • EQUIPO TÉCNICO DE SISTEMAS Y ARQUITECTURA DIGITAL"):
        # Top banner category
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_BLUE_MIGATO

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(10.5), Inches(0.75))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY_DARK

        # Logo top right
        if LOGO_PATH.exists():
            slide.shapes.add_picture(str(LOGO_PATH), Inches(11.8), Inches(0.4), Inches(0.75), Inches(0.75))

        # Bottom subtle separator
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_NAVY_CORP
        line.line.fill.background()

    def add_speaker_note(slide, note_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = note_text

    # ==========================================
    # SLIDE 1: PORTADA INSTITUCIONAL / BUSINESS
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    # Logo Centrado
    if LOGO_PATH.exists():
        s1.shapes.add_picture(str(LOGO_PATH), Inches(5.9), Inches(0.9), Inches(1.5), Inches(1.5))

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.333), Inches(4.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "REPÚBLICA BOLIVARIANA DE VENEZUELA • ESTADO MONAGAS"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_MIGATO

    p = tf1.add_paragraph()
    p.text = "MOVIMIENTO INDEPENDIENTE GANAMOS TODOS (MIGATO)"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "ARQUITECTURA DE GESTIÓN TERRITORIAL Y PROPUESTA DE FACTIBILIDAD TÉCNICA"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK
    p.space_after = Pt(10)

    p = tf1.add_paragraph()
    p.text = "Evaluación Integral de los Cinco Módulos de Mando, Servidor Cloud VPS Privado y Presupuesto Operativo"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_after = Pt(22)

    p = tf1.add_paragraph()
    p.text = "Presentado por el Equipo Técnico de Sistemas y Arquitectura Digital\nA la Dirección General de MIGATO • Maturín, Septiembre 2026"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_BODY

    add_speaker_note(s1, "GUION: Estimados miembros de la Dirección General: les presentamos el dictamen técnico y la propuesta de factibilidad para la puesta en marcha de la Plataforma Territorial MIGATO. El desarrollo del software ha sido completado al 100% por nuestro equipo sin costo de honorarios para la organización. El objetivo de esta sesión ejecutiva es someter a su consideración la aprobación de un presupuesto operativo mínimo de $24.50 mensuales para la infraestructura cloud del servidor privado.")

    # ==========================================
    # SLIDE 2: MARCO ESTRATÉGICO
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "Marco Estratégico: Superación del Centralismo Burocrático")

    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_WHITE
    c1.line.color.rgb = COLOR_BORDER

    tb_c1 = s2.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "⚠️ Limitaciones del Esquema Tradicional"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_DANGER_DARK
    p.space_after = Pt(12)

    bullets_c1 = [
        "Precarización de servicios esenciales: fallas sostenidas en la red eléctrica, suministro de agua y vialidad.",
        "Gestión sin métricas objetivas: ausencia de levantamiento técnico y decisiones basadas en criterios clientelares.",
        "Riesgos en canales informales: la dispersión en grupos de mensajería compromete la confidencialidad y dificulta la consolidación.",
        "Vulnerabilidad operativa: exposición de los activistas comunitarios ante revisiones no autorizadas de dispositivos en campo."
    ]
    for b in bullets_c1:
        p = tf_c1.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_WHITE
    c2.line.color.rgb = COLOR_NAVY_CORP

    tb_c2 = s2.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "🏛️ La Propuesta Gerencial de MIGATO"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(12)

    bullets_c2 = [
        "Estándares corporativos de primer nivel: implementación de orden, método y rigurosidad técnica en la gestión pública.",
        "Decisiones con respaldo empírico: información georreferenciada y métricas confiables de las comunidades en tiempo real.",
        "Comunicación estructurada y segura: canales compartimentados por parroquia con trazabilidad institucional.",
        "Soberanía de la información: bases de datos resguardadas en un servidor cloud dedicado bajo control exclusivo del partido."
    ]
    for b in bullets_c2:
        p = tf_c2.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    add_speaker_note(s2, "GUION: La gestión pública en Monagas enfrenta una severa ineficiencia burocrática que ha deteriorado los servicios básicos y la confianza ciudadana. En el ámbito organizativo, operar mediante canales informales de mensajería genera dispersión, pérdida de datos y riesgos de seguridad para los equipos de campo. Frente a este panorama, MIGATO propone un modelo de gestión basado en la competencia técnica, la transparencia administrativa y la toma de decisiones fundamentada en indicadores reales y auditables.")

    # ==========================================
    # SLIDE 3: SOLUCIÓN EN 5 MÓDULOS
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "Solución Integral: Plataforma Territorial en 5 Módulos")

    modulos = [
        ("Módulo 1: Despacho Celular Seguro", "/despacho/", "Enlaces cifrados por token web temporal directo al coordinador. Sin contraseñas complejas y con compartimentación territorial."),
        ("Módulo 2: Censo Territorial y Privacidad", "/carga/", "Protocolo de Cero Residuos en Dispositivo: registra variables cuantitativas de servicios sin almacenar datos personales ni cédulas."),
        ("Módulo 3: Padrón Electoral y Actas", "/centros-maturin/", "Catálogo de 175 centros y 361 mesas de Maturín. Organización de testigos y respaldo fotográfico de actas y código QR para auditoría interna."),
        ("Módulo 4: Cartografía Tridimensional", "/earth-monagas/", "Consola 3D de vialidad, relieve y servicios con filtro de focalización territorial para mesas de planificación estratégica."),
        ("Módulo 5: Diagnóstico de Salud Regional", "/salud-monagas/", "Supervisión técnica de 84 centros asistenciales: plantas eléctricas, agua, quirófanos y base fehaciente para planes de gobierno.")
    ]

    for idx, (m_title, m_route, m_desc) in enumerate(modulos):
        y_pos = Inches(1.8 + idx * 0.98)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(0.88))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s3.shapes.add_textbox(Inches(1.0), y_pos + Inches(0.08), Inches(11.3), Inches(0.72))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{m_title}   "
        p.font.name = "Arial"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY_CORP

        p2 = tf.add_paragraph()
        p2.text = f"Consola: {m_route} • {m_desc}"
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_BODY

    add_speaker_note(s3, "GUION: La plataforma constituye una solución tecnológica completa y operativa que articula cinco componentes complementarios: despacho parroquial por enlace web seguro, censo comunitario con resguardo de privacidad, seguimiento del padrón electoral con archivo de actas, consola cartográfica 3D y supervisión técnica de la red regional de salud. Todo el sistema opera bajo una arquitectura web unificada, accesible y sin intermediarios.")

    # ==========================================
    # SLIDE 4: MÓDULOS 1 Y 2
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "Módulos 1 y 2: Coordinación Parroquial y Protección de Datos")

    c1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_WHITE
    c1.line.color.rgb = COLOR_BORDER

    tb1 = s4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "📲 Módulo 1: Despacho Celular Seguro"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(12)

    bullets_s4_1 = [
        "Acceso Instantáneo por Token Web: enlaces temporales cifrados enviados directamente al coordinador; ingreso inmediato sin necesidad de memorizar contraseñas.",
        "Compartimentación Territorial: cada enlace opera exclusivamente en su área asignada, garantizando estricta segmentación de responsabilidades.",
        "Revocación Inmediata de Credenciales: ante cualquier relevo de personal, la credencial se desactiva desde la consola central sin interrumpir el servicio.",
        "Compatibilidad Universal: funciona directamente en navegadores móviles estándar sin requerir instalación de aplicaciones adicionales."
    ]
    for b in bullets_s4_1:
        p = tf1.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    c2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_WHITE
    c2.line.color.rgb = COLOR_NAVY_CORP

    tb2 = s4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🛡️ Módulo 2: Censo Territorial y Privacidad"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(12)

    bullets_s4_2 = [
        "Protección de Identidad Ciudadana: por directriz institucional, el sistema no solicita nombres, cédulas ni números telefónicos de los vecinos censados.",
        "Diagnóstico de Servicios Básicos: levantamiento riguroso de familias, viviendas y fallas críticas en suministro eléctrico, agua potable y vialidad.",
        "Protocolo Zero-Byte Storage: al presionar enviar, los datos se transmiten al servidor central vía HTTPS y el formulario se restablece. No queda caché en el celular.",
        "Tranquilidad Operativa: blindaje total del activista frente a revisiones no autorizadas de dispositivos en la vía pública."
    ]
    for b in bullets_s4_2:
        p = tf2.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    add_speaker_note(s4, "GUION: En el trabajo de campo, la seguridad de nuestros coordinadores y la privacidad ciudadana son prioritarias. Por ello, el Módulo 1 elimina contraseñas mediante tokens seguros que dan acceso exclusivo a la parroquia asignada. En el Módulo 2, el censo se enfoca en necesidades cuantitativas (familias, viviendas, transformadores) y, por principio ético, no recopila nombres ni cédulas. Al enviar el formulario, los datos viajan cifrados y el teléfono queda completamente limpio.")

    # ==========================================
    # SLIDE 5: MÓDULO 3 (AUDITORÍA ELECTORAL)
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Módulo 3: Padrón Electoral, Centros de Votación y Resguardo de Actas")

    cols_s5 = [
        ("🗺️ Catálogo de Centros", [
            "Georreferenciación completa de 175 centros y 361 mesas electorales del Municipio Maturín.",
            "Identificación prioritaria de los 40 centros estratégicos que concentran más del 60% de los electores.",
            "Mapeo de accesos viales y distancias logísticas para optimizar la movilización."
        ]),
        ("👥 Estructura de Testigos", [
            "Registro y acreditación organizada de testigos principales y suplentes por mesa electoral.",
            "Segmentación de labores de asistencia ciudadana para adultos mayores y votantes jóvenes.",
            "Canal de reporte de incidencias y apertura de mesas en comunicación con la Sala Situacional."
        ]),
        ("📄 Archivo de Evidencia", [
            "Carga ordenada de resultados numéricos por mesa emitida por los testigos del partido.",
            "Almacenamiento de respaldo digital del acta física y comprobante oficial.",
            "Registro fotográfico del código QR para auditoría y cotejo interno de la organización.",
            "Sin intervenir ni suplantar las competencias del CNE."
        ]),
    ]

    for idx, (col_title, col_bullets) in enumerate(cols_s5):
        x_pos = Inches(0.8 + idx * 4.0)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.8), Inches(3.733), Inches(5.0))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s5.shapes.add_textbox(x_pos + Inches(0.2), Inches(2.0), Inches(3.333), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = col_title
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY_CORP
        p.space_after = Pt(12)

        for b in col_bullets:
            p = tf.add_paragraph()
            p.text = "▪ " + b
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT_BODY
            p.space_after = Pt(8)

    add_speaker_note(s5, "GUION: El Módulo 3 está concebido como una herramienta de archivo y auditoría interna para la organización, respetando plenamente el marco institucional. Mapea la totalidad de los 175 centros y 361 mesas de Maturín, identificando los 40 centros que concentran más del 60% del padrón. Facilita la asignación de testigos y permite cargar los resultados con respaldo fotográfico del acta física y comprobante impreso con su código QR, asegurando un expediente fehaciente para nuestro cotejo interno.")

    # ==========================================
    # SLIDE 6: MÓDULOS 4 Y 5
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Módulos 4 y 5: Cartografía Tridimensional y Auditoría Asistencial")

    c1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_WHITE
    c1.line.color.rgb = COLOR_BORDER

    tb1 = s6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🌐 Cartografía 3D de Infraestructura (/earth-monagas/)"
    p.font.name = "Arial"
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(12)

    bullets_s6_1 = [
        "Análisis Geoespacial Interactivo: representación tridimensional de relieve, vialidad, comunidades y equipamiento de servicios en la entidad.",
        "Herramienta de Focalización Territorial: filtro visual que atenúa suavemente las zonas periféricas para concentrar el análisis en la parroquia de interés.",
        "Planificación de Rutas y Despliegue: medición precisa de distancias y estimación de tiempos de traslado para jornadas institucionales de trabajo.",
        "Optimizado para Sala de Mando: interfaz adaptada para proyección en pantallas corporativas de alta definición."
    ]
    for b in bullets_s6_1:
        p = tf1.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    c2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_WHITE
    c2.line.color.rgb = COLOR_NAVY_CORP

    tb2 = s6.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🏥 Diagnóstico de la Red de Salud (/salud-monagas/)"
    p.font.name = "Arial"
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(12)

    bullets_s6_2 = [
        "Supervisión Integral de 84 Centros: monitoreo técnico de hospitales universitarios, ambulatorios urbanos, rurales y centros integrales.",
        "Estado de Servicios Críticos: operatividad de generadores eléctricos de emergencia, suministro continuo de agua y salas quirúrgicas.",
        "Red de Referencia y Traslado: registro de las rutas de derivación de pacientes cuando la capacidad de resolución local se ve superada.",
        "Fundamento para Políticas Públicas: base de datos auditable que sustenta propuestas de rehabilitación sanitaria con criterios de ingeniería."
    ]
    for b in bullets_s6_2:
        p = tf2.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    add_speaker_note(s6, "GUION: El Módulo 4 ofrece a la Dirección una consola geoespacial tridimensional para evaluar la topografía, la red de carreteras y los servicios públicos. Incorpora un filtro de focalización territorial que atenúa las áreas circundantes para centrar el debate ejecutivo en la parroquia evaluada. Por su parte, el Módulo 5 audita técnicamente 84 centros de salud de Monagas, registrando el estado de plantas eléctricas y quirófanos, lo que provee el fundamento técnico indispensable para nuestras propuestas de gobierno.")

    # ==========================================
    # SLIDE 7: INFRAESTRUCTURA CLOUD Y SEGURIDAD
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Seguridad de la Información y Servidor Cloud VPS Dedicado")

    boxes_s7 = [
        ("🔒 Cifrado SSL/TLS de 256 Bits", "Encriptación completa de extremo a extremo en todas las sesiones web, impidiendo la intercepción de tráfico o fuga de datos en redes públicas o conexiones móviles."),
        ("🛡️ Auditoría de Vulnerabilidades", "Cierre estricto de puertos no esenciales, filtrado de conexiones mediante reglas iptables y mitigación automatizada de intentos de intrusión repetitivos (Fail2ban)."),
        ("⚡ Pruebas de Estrés y Concurrencia", "Simulaciones de alta demanda de tráfico para certificar que el servidor mantenga tiempos de respuesta óptimos (menores a 1 segundo) durante jornadas intensas de carga."),
        ("💾 Respaldos Automatizados Fuera de Línea", "Generación periódica de copias de seguridad encriptadas fuera del servidor principal para asegurar la continuidad del servicio ante cualquier eventualidad técnica.")
    ]

    for idx, (b_title, b_desc) in enumerate(boxes_s7):
        row = idx // 2
        col = idx % 2
        x_pos = Inches(0.8 + col * 6.0)
        y_pos = Inches(1.8 + row * 2.5)

        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, y_pos, Inches(5.7), Inches(2.25))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s7.shapes.add_textbox(x_pos + Inches(0.25), y_pos + Inches(0.2), Inches(5.2), Inches(1.85))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = b_title
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY_CORP
        p.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = b_desc
        p2.font.name = "Arial"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = COLOR_TEXT_BODY

    add_speaker_note(s7, "GUION: Una solución estratégica de esta magnitud requiere una infraestructura profesional. A diferencia de los alojamientos compartidos que sufren caídas imprevistas y carecen de aislamiento, un Servidor Virtual Privado dedicado nos otorga independencia, procesamiento exclusivo y respaldo continuo. Implementamos cifrado SSL de 256 bits, reglas de cortafuegos y pruebas de estrés para asegurar que la plataforma mantenga tiempos de respuesta inferiores a un segundo durante picos de concurrencia.")

    # ==========================================
    # SLIDE 8: ESTUDIO ECONÓMICO
    # ==========================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "Estudio Económico y Presupuesto de Inversión: Fase 1")

    # Banner KPI
    c_banner = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.733), Inches(0.85))
    c_banner.fill.solid()
    c_banner.fill.fore_color.rgb = COLOR_BLUE_LIGHT
    c_banner.line.color.rgb = COLOR_BLUE_MIGATO

    tb_b = s8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(0.75))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "Presupuesto Solicitado: ~$24.50 USD / mes   |   $294.00 USD / AÑO TOTAL"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP

    p2 = tf_b.add_paragraph()
    p2.text = "Inversión exclusiva en infraestructura de nube • Honorarios de desarrollo provistos como aporte técnico ($0.00)"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Arial"
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_TEXT_BODY

    # Table
    rows = 6
    cols = 4
    left = Inches(0.8)
    top = Inches(2.8)
    width = Inches(11.733)
    height = Inches(3.8)

    table_shape = s8.shapes.add_table(rows, cols, left, top, width, height)
    t = table_shape.table

    t.columns[0].width = Inches(3.2)
    t.columns[1].width = Inches(5.133)
    t.columns[2].width = Inches(1.7)
    t.columns[3].width = Inches(1.7)

    headers = ["Componente de Infraestructura", "Especificación Técnica Mínima", "Costo Mensual", "Costo Anual"]
    for i, h in enumerate(headers):
        cell = t.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY_CORP
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_WHITE
        if i >= 2:
            p.alignment = PP_ALIGN.CENTER

    data = [
        ("Servidor Cloud VPS Dedicado", "4 vCPU / 8 GB RAM / 100 GB SSD NVMe / Red 1 Gbps / Protección contra ataques", "$20.00 USD", "$240.00 USD"),
        ("Dominio Institucional Oficial", "Registro anual de dominio web corporativo + Certificado SSL Wildcard (256 bits)", "$1.50 USD", "$18.00 USD"),
        ("IP Pública Fija Dedicada", "Dirección IP estática exclusiva para filtrado de seguridad en consola de mando", "$3.00 USD", "$36.00 USD"),
        ("Software de Servidor y Código", "Entorno Linux Ubuntu Server, Nginx, PostgreSQL, Fail2ban y módulos FOSS", "$0.00 USD", "$0.00 USD"),
        ("TOTAL MÍNIMO SOLICITADO (FASE 1):", "Infraestructura cloud dedicada y certificada para las pruebas piloto", "~$24.50 USD", "$294.00 USD")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        is_total = (row_idx == 5)
        for col_idx, text in enumerate(row_data):
            cell = t.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BLUE_LIGHT if is_total else (COLOR_WHITE if row_idx % 2 == 1 else COLOR_CARD_SUBTLE)
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Arial"
            p.font.size = Pt(11 if not is_total else 11.5)
            p.font.bold = is_total or (col_idx == 0)
            p.font.color.rgb = COLOR_NAVY_CORP if is_total else COLOR_TEXT_MAIN
            if col_idx >= 2:
                p.alignment = PP_ALIGN.CENTER
                if is_total:
                    p.font.color.rgb = COLOR_NAVY_CORP

    add_speaker_note(s8, "GUION: En esta lámina presentamos el análisis económico con estricta transparencia. El costo del desarrollo de software y arquitectura representa cero bolívares y cero dólares para la organización, habiendo sido aportado íntegramente por el equipo técnico. La inversión solicitada se limita exclusivamente a los costos directos de la nube: $20 mensuales por el VPS dedicado, $1.50 por el dominio y certificados SSL, y $3 por la IP fija. El total asciende a 24 dólares y medio mensuales, o 294 dólares anuales.")

    # ==========================================
    # SLIDE 9: PLAN DE DESPLIEGUE ESCALONADO
    # ==========================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "Plan de Despliegue Escalonado por Fases")

    c1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_WHITE
    c1.line.color.rgb = COLOR_SUCCESS_DARK

    tb1 = s9.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🚀 FASE 1: Inmediata (Objeto de esta Solicitud)"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS_DARK
    p.space_after = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Inversión requerida: $294.00 USD anuales"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.italic = True
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_after = Pt(12)

    f1_bullets = [
        "Contratación del servidor cloud VPS y vinculación del dominio institucional.",
        "Instalación de la arquitectura de seguridad, reglas de firewall y certificados SSL de 256 bits.",
        "Ejecución de pruebas de carga, concurrencia y auditoría técnica de puertos.",
        "Validación operativa en 3 parroquias piloto de Maturín: Las Cocuizas, San Simón y Alto de Los Godos con coordinadores reales."
    ]
    for b in f1_bullets:
        p = tf1.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_BODY
        p.space_after = Pt(8)

    c2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_WHITE
    c2.line.color.rgb = COLOR_BORDER

    tb2 = s9.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🏢 FASE 2: Consolidación de Sala Situacional Física"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK
    p.space_after = Pt(6)

    p = tf2.add_paragraph()
    p.text = "Planificación diferida sujeta a disponibilidad de recursos:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.italic = True
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_after = Pt(12)

    f2_bullets = [
        "Acondicionamiento físico del centro de seguimiento y monitoreo territorial.",
        "Dotación de equipos de computación dedicados para operadores de consola.",
        "Instalación de pantallas de visualización general para seguimiento en tiempo real.",
        "Sistemas de respaldo eléctrico (inversores / UPS) para garantizar continuidad operativa ante fallas del servicio eléctrico regional."
    ]
    for b in f2_bullets:
        p = tf2.add_paragraph()
        p.text = "▪ " + b
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    add_speaker_note(s9, "GUION: Con un criterio gerencial responsable, proponemos una implementación escalonada en dos etapas. La presente solicitud corresponde exclusivamente a la Fase 1, con una inversión única de 294 dólares anuales para encender el servidor y validar la plataforma en tres parroquias piloto de Maturín: Las Cocuizas, San Simón y Alto de Los Godos. La Fase 2, relativa al equipamiento físico de la Sala Situacional, se evaluará posteriormente según la disponibilidad y planificación del partido.")

    # ==========================================
    # SLIDE 10: CONCLUSIÓN Y RECOMENDACIÓN FINAL
    # ==========================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10)
    add_header(s10, "Conclusión y Recomendación Final del Equipo Técnico")

    c_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(10.933), Inches(5.0))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = COLOR_WHITE
    c_box.line.color.rgb = COLOR_NAVY_CORP

    tb_box = s10.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.5))
    tf_box = tb_box.text_frame
    tf_box.word_wrap = True

    p = tf_box.paragraphs[0]
    p.text = "✅ Dictamen Favorable y Factibilidad Técnica"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_CORP
    p.space_after = Pt(14)

    conclusions = [
        ("Madurez Funcional Comprobada:", "Los cinco módulos se encuentran programados y evaluados en entornos de prueba, garantizando plena operatividad sin riesgos de desarrollo."),
        ("Seguridad y Protección Institucional:", "El esquema de tokens web efímeros y formularios sin almacenamiento residual en dispositivos brinda tranquilidad y resguardo absoluto a la militancia."),
        ("Alta Rentabilidad de la Inversión:", "El monto solicitado ($294.00 USD anuales) cubre exclusivamente costos directos de servidor cloud, representando un valor mínimo frente a su alto impacto estratégico.")
    ]
    for title, desc in conclusions:
        p = tf_box.add_paragraph()
        p.text = f"▪ {title} "
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY_DARK
        p.space_after = Pt(2)

        p_desc = tf_box.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = COLOR_TEXT_BODY
        p_desc.space_after = Pt(8)

    p_rec = tf_box.add_paragraph()
    p_rec.text = "RECOMENDACIÓN FINAL: Aprobar el presupuesto operativo de $294.00 USD anuales para proceder de inmediato con la contratación del servidor dedicado y el encendido formal de la Fase 1 del proyecto."
    p_rec.font.name = "Arial"
    p_rec.font.size = Pt(12.5)
    p_rec.font.bold = True
    p_rec.font.color.rgb = COLOR_NAVY_CORP
    p_rec.space_before = Pt(8)

    add_speaker_note(s10, "GUION: Para concluir, reiteramos que la herramienta se encuentra completamente terminada, probada y lista para entrar en operación. Ofrece una alta solvencia funcional y un protocolo estricto de resguardo para la militancia. Con una inversión mínima de 294 dólares anuales, MIGATO se dota de una plataforma corporativa soberana. El Equipo Técnico recomienda formalmente la aprobación de este presupuesto para iniciar el encendido de inmediato. Quedamos a su disposición.")

    prs.save(str(filename))
    print(f"[✓] Presentación PowerPoint Corporativa generada exitosamente: {filename} ({os.path.getsize(filename)} bytes)")

if __name__ == "__main__":
    out_file = PROJECT_ROOT / "PRESENTACION_EJECUTIVA_MIGATO_V3.pptx"
    create_deck(out_file)
